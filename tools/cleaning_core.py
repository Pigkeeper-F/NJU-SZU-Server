"""Deterministic table inspection, frozen-script execution and independent checks.

No model calls. CLI commands: profile, run, verify. The same file is included
in reproduction bundles; input records are never rewritten in place.
"""
from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import importlib.util
import io
import json
import math
import os
import platform
import re
import sys
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path

VERSION = 1
FORMATS = {".csv", ".tsv", ".xlsx", ".xls"}
MAX_ROWS = 500000
MAX_CELLS = 5000000


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n"


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".tmp")
    temp.write_text(dumps(value), encoding="utf-8")
    os.replace(temp, path)


def safe_file(root, name):
    if not isinstance(name, str) or not name or Path(name).name != name or "/" in name or "\\" in name or name.startswith("."):
        raise ValueError("无效输入文件名")
    root = Path(root).resolve()
    target = root / name
    if target.is_symlink() or not target.is_file() or target.resolve().parent != root:
        raise ValueError("输入文件不存在或不是当前目录内的普通文件：" + name)
    return target


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("表格包含非有限数值")
    return str(value)


def runtime():
    from importlib.metadata import PackageNotFoundError, version
    packages = {}
    for name in ("openpyxl", "xlrd"):
        try:
            packages[name] = version(name)
        except PackageNotFoundError:
            packages[name] = "unknown"
    return {"python": platform.python_version(), "platform": platform.system(),
            "machine": platform.machine(), "packages": packages}


def _excel(path, chosen):
    tables, sheets = [], []
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        # Normal worksheets expose merged ranges as well as formula cells.
        values = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=False)
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=False)
        try:
            sheets = values.sheetnames
            wanted = sheets if chosen is None else chosen
            if any(name not in sheets for name in wanted):
                raise ValueError("指定工作表不存在：" + path.name)
            for name in wanted:
                sheet, source = values[name], formulas[name]
                if sheet.max_row > MAX_ROWS + 1 or sheet.max_row * sheet.max_column > MAX_CELLS:
                    raise ValueError("Excel工作表超过本轮检查上限；请先限定数据范围")
                rows = [[cell_text(c.value) for c in row] for row in sheet.iter_rows()]
                missing_cache = [cell.coordinate for row in source.iter_rows() for cell in row
                                 if cell.data_type == "f" and sheet[cell.coordinate].value is None]
                warnings = []
                if source.merged_cells.ranges:
                    warnings.append({"code": "merged_cells", "blocking": True,
                                     "message": "存在合并单元格，请明确表头/拆分规则", "cells": [str(r) for r in source.merged_cells.ranges][:30]})
                if missing_cache:
                    warnings.append({"code": "formula_cache_missing", "blocking": True,
                                     "message": "公式没有可读取的缓存值；请提供已保存计算结果或明确处理方法", "cells": missing_cache[:30], "count": len(missing_cache)})
                errors = [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "e"]
                if errors:
                    warnings.append({"code": "excel_error_cells", "blocking": True, "message": "存在Excel错误值，请明确处理方法", "cells": errors[:30], "count": len(errors)})
                formats = sorted({c.number_format for row in source.iter_rows() for c in row if c.value is not None and c.number_format != "General"})
                tables.append((name, rows, warnings, {"formula_policy": "cached_values_only", "number_formats": formats[:30]}))
        finally:
            values.close()
            formulas.close()
    else:
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
        try:
            sheets = book.sheet_names()
            wanted = sheets if chosen is None else chosen
            if any(name not in sheets for name in wanted):
                raise ValueError("指定工作表不存在：" + path.name)
            for name in wanted:
                sheet = book.sheet_by_name(name)
                if sheet.nrows > MAX_ROWS + 1 or sheet.nrows * sheet.ncols > MAX_CELLS:
                    raise ValueError("Excel工作表超过本轮检查上限；请先限定数据范围")
                errors, dates, rows = [], [], []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        value, kind = sheet.cell_value(r, c), sheet.cell_type(r, c)
                        coordinate = "R" + str(r + 1) + "C" + str(c + 1)
                        if kind == xlrd.XL_CELL_ERROR:
                            errors.append(coordinate)
                            value = xlrd.error_text_from_code.get(value, "unknown_excel_error")
                        elif kind == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif kind == xlrd.XL_CELL_DATE:
                            dates.append(coordinate)
                        row.append(cell_text(value))
                    rows.append(row)
                issues = [{"code": "xls_formula_provenance_unknown", "blocking": False,
                           "message": "XLS读取保存值；此格式下不能可靠区分公式及缺失缓存，公式来源状态为unknown"}]
                if sheet.merged_cells:
                    issues.append({"code": "merged_cells", "blocking": True, "message": "存在合并单元格，请明确表头/拆分规则"})
                if errors:
                    issues.append({"code": "excel_error_cells", "blocking": True, "message": "存在Excel错误值，请明确处理方法", "cells": errors[:30], "count": len(errors)})
                tables.append((name, rows, issues, {"formula_policy": "saved_values; formula provenance unknown",
                                                   "date_cells": dates[:30], "date_cell_count": len(dates), "date_mode": book.datemode}))
        finally:
            book.release_resources()
    return tables, sheets


def load_tables(input_dir, selection):
    if not isinstance(selection, list) or not 1 <= len(selection) <= 30:
        raise ValueError("请选择1至30个输入文件")
    tables, inputs = [], []
    seen = set()
    for entry in selection:
        if not isinstance(entry, dict) or entry.get("name") in seen:
            raise ValueError("输入选择无效或重复")
        name = entry.get("name")
        path = safe_file(input_dir, name)
        seen.add(name)
        if path.suffix.lower() not in FORMATS:
            raise ValueError("清洗检查支持CSV、TSV、XLSX和XLS")
        if entry.get("sha256") and sha(path) != entry["sha256"]:
            raise ValueError("输入指纹已变化：" + name)
        descriptor = {"name": name, "sha256": sha(path), "size": path.stat().st_size,
                      "encoding": entry.get("encoding", "utf-8-sig"), "sheets": entry.get("sheets")}
        if descriptor["sheets"] is not None and (not isinstance(descriptor["sheets"], list) or not descriptor["sheets"] or not all(isinstance(s, str) for s in descriptor["sheets"])):
            raise ValueError("工作表选择不能为空或重复")
        if descriptor["sheets"] and len(set(descriptor["sheets"])) != len(descriptor["sheets"]):
            raise ValueError("工作表选择重复")
        if path.suffix.lower() in {".csv", ".tsv"}:
            encoding = descriptor["encoding"]
            if encoding not in {"utf-8-sig", "utf-16", "gb18030"}:
                raise ValueError("请选择支持的文本编码")
            with path.open(encoding=encoding, newline="") as stream:
                reader = csv.reader(stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ",", strict=True)
                rows = []
                for row in reader:
                    if len(rows) > MAX_ROWS or (len(rows) + 1) * max(len(row), 1) > MAX_CELLS:
                        raise ValueError("文件超过本轮完整检查上限；请先限定数据范围")
                    rows.append(row)
            sources = [(None, rows, [], {"encoding": encoding})]
        else:
            sources, available = _excel(path, descriptor["sheets"])
            descriptor["available_sheets"] = available
        inputs.append(descriptor)
        for sheet, rows, issues, metadata in sources:
            columns = rows[0] if rows else []
            data = rows[1:] if rows else []
            if not columns or any(not str(c).strip() for c in columns) or len(set(columns)) != len(columns):
                issues.append({"code": "ambiguous_header", "blocking": True, "message": "表头为空或重复，请明确表头位置和字段名"})
            uneven = [i + 2 for i, row in enumerate(data) if len(row) != len(columns)]
            if uneven:
                issues.append({"code": "uneven_columns", "blocking": True, "message": "记录列数与表头不一致", "rows": uneven[:30], "count": len(uneven)})
            tables.append({"id": "T" + str(len(tables) + 1), "file": name, "sheet": sheet,
                           "columns": columns, "rows": data, "issues": issues, "metadata": metadata})
    return tables, inputs


def describe(table):
    rows, columns = table["rows"], table["columns"]
    fields = []
    for index, name in enumerate(columns):
        values = [row[index] if index < len(row) else "" for row in rows]
        counts = collections.Counter(values)
        numbers = []
        for record, value in enumerate(values, 1):
            if not value.strip():
                continue
            try:
                number = Decimal(value)
                if number.is_finite():
                    numbers.append((record, number))
            except InvalidOperation:
                pass
        fields.append({"name": name, "empty": sum(not v.strip() for v in values),
                       "surrounding_whitespace": sum(v != v.strip() for v in values),
                       "distinct": len(counts), "values": dict(counts) if len(counts) <= 30 else None,
                       "numeric_count": len(numbers), "zero_count": sum(n == 0 for _, n in numbers),
                       "negative_count": sum(n < 0 for _, n in numbers),
                       "min": str(min(n for _, n in numbers)) if numbers else None,
                       "max": str(max(n for _, n in numbers)) if numbers else None})
    return {k: v for k, v in table.items() if k != "rows"} | {
        "row_count": len(rows), "duplicate_rows": len(rows) - len({tuple(r) for r in rows}),
        "blank_records": sum(not any(v.strip() for v in row) for row in rows),
        "fields": fields, "sample": rows[:5]}


def profile(input_dir, selection):
    tables, inputs = load_tables(input_dir, selection)
    result = {"schema": VERSION, "inputs": inputs, "tables": [describe(t) for t in tables],
              "complete_scan": True, "row_count": sum(len(t["rows"]) for t in tables)}
    result["blocked"] = any(issue.get("blocking") for table in tables for issue in table["issues"])
    return result


def read_bundle(bundle):
    bundle = Path(bundle)
    plan = json.loads(safe_file(bundle, "cleaning-plan.json").read_text(encoding="utf-8"))
    if plan.get("schema") != VERSION or not plan.get("inputs") or not plan.get("rules", {}).get("text"):
        raise ValueError("清洗任务清单无效")
    if sha(safe_file(bundle, "clean.py")) != plan.get("script_sha256"):
        raise ValueError("冻结的清洗脚本已变化")
    if sha(safe_file(bundle, "cleaning_core.py")) != plan.get("engine_sha256"):
        raise ValueError("冻结的清洗内核已变化")
    return plan


def execute(bundle, input_dir, output_dir):
    bundle, output = Path(bundle), Path(output_dir)
    plan = read_bundle(bundle)
    tables, _ = load_tables(input_dir, plan["inputs"])
    if any(issue.get("blocking") for t in tables for issue in t["issues"]):
        raise ValueError("输入结构存在待确认问题，未执行清洗")
    payload = []
    for table in tables:
        payload.append({"id": table["id"], "file": table["file"], "sheet": table["sheet"],
                        "columns": table["columns"], "rows": [dict(zip(table["columns"], row)) for row in table["rows"]],
                        "source_ids": [table["id"] + ":R" + str(i + 1) for i in range(len(table["rows"]))]})
    output.mkdir(parents=True, exist_ok=True)
    if (output / "cleaning-result.json").exists():
        raise ValueError("结果已存在；请使用新的输出目录")
    spec = importlib.util.spec_from_file_location("frozen_project_cleaner", bundle / "clean.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    # The project function returns data; fixed code owns serialization.
    result = module.clean(payload, {"rules": plan["rules"]})
    if isinstance(result, dict) and result.get("questions"):
        questions = result["questions"]
        if not isinstance(questions, list) or len(questions) > 20 or not all(isinstance(q, str) and 0 < len(q) <= 1000 for q in questions):
            raise ValueError("脚本返回的待确认问题无效")
        write_json(output / "needs-attention.json", {"status": "needs_attention", "questions": questions})
        return {"questions": questions}
    if not isinstance(result, dict) or not isinstance(result.get("tables"), list) or not result["tables"]:
        raise ValueError("脚本必须返回非空tables列表")
    if len(result["tables"]) > 30:
        raise ValueError("输出表过多")
    for index, table in enumerate(result["tables"], 1):
        columns, rows, lineage = table.get("columns"), table.get("rows"), table.get("lineage")
        if not isinstance(columns, list) or not columns or not all(isinstance(c, str) and c for c in columns) or len(set(columns)) != len(columns):
            raise ValueError("输出字段为空或重复")
        if not isinstance(rows, list) or not isinstance(lineage, list) or len(rows) != len(lineage):
            raise ValueError("输出必须提供逐行来源lineage")
        if len(rows) > MAX_ROWS:
            raise ValueError("输出超过记录上限")
        if any(not isinstance(row, dict) or set(row) != set(columns) for row in rows):
            raise ValueError("输出行字段不一致")
        table["rows"] = [{c: cell_text(row[c]) for c in columns} for row in rows]
        table["file"] = "table-" + str(index).zfill(2) + ".csv"
        with (output / table["file"]).open("x", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, lineterminator="\n")
            writer.writeheader()
            writer.writerows(table["rows"])
    write_json(output / "cleaning-result.json", result)
    return {"tables": len(result["tables"]), "rows": sum(len(t["rows"]) for t in result["tables"])}


def verify(bundle, input_dir, output_dir):
    plan = read_bundle(bundle)
    tables, _ = load_tables(input_dir, plan["inputs"])
    output = Path(output_dir)
    result_path = safe_file(output, "cleaning-result.json")
    if result_path.stat().st_size > 128 * 1024 * 1024:
        raise ValueError("结果清单超过本轮核验大小上限")
    result = json.loads(result_path.read_text(encoding="utf-8"))
    source = {t["id"] + ":R" + str(i + 1): dict(zip(t["columns"], row)) for t in tables for i, row in enumerate(t["rows"])}
    checks, used = [], set()
    def check(name, passed, detail=""):
        checks.append({"id": name, "status": "passed" if passed else "failed", "detail": detail})
    output_tables = result.get("tables")
    if not isinstance(output_tables, list) or not output_tables:
        raise ValueError("缺少清洗结果表")
    for index, table in enumerate(output_tables, 1):
        rows, columns, lineage = table["rows"], table["columns"], table["lineage"]
        expected = "table-" + str(index).zfill(2) + ".csv"
        if table.get("file") != expected:
            raise ValueError("结果文件名称与固定序列不符")
        with safe_file(output, expected).open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            saved = list(reader)
            check("output_" + str(index) + "_readback", reader.fieldnames == columns and saved == rows)
        valid = len(rows) == len(lineage) and all(isinstance(ids, list) and ids and all(isinstance(s, str) and s in source for s in ids) for ids in lineage)
        check("output_" + str(index) + "_lineage", valid)
        if valid:
            used.update(s for ids in lineage for s in ids)
    dispositions = result.get("dispositions", {})
    valid_dispositions = isinstance(dispositions, dict) and all(k in source and isinstance(v, str) and v.strip() for k, v in dispositions.items())
    check("dispositions_valid", valid_dispositions)
    explained = set(dispositions) if valid_dispositions else set()
    check("source_records_accounted", used | explained == set(source), "每条输入必须进入输出或有明确去向；这不代表删除理由已获科学验收")
    check("used_records_not_disposed", not used & explained)
    check("input_hashes_unchanged", all(sha(safe_file(input_dir, e["name"])) == e["sha256"] for e in plan["inputs"]))
    # Only fixed check operators are accepted; evidence is tied to confirmed text.
    for index, rule in enumerate(plan.get("checks", []), 1):
        basis = rule.get("basis", "")
        if not basis or basis not in plan["rules"]["text"]:
            raise ValueError("项目检查缺少已确认规则的原句依据")
        target = output_tables[int(rule.get("table", 1)) - 1]
        op, fields = rule.get("op"), rule.get("fields", [])
        values = target["rows"]
        if op == "row_count":
            passed = len(values) == rule["expected"]
        elif op == "unique":
            keys = [tuple(row[f] for f in fields) for row in values]
            passed = len(keys) == len(set(keys))
        elif op == "not_null":
            passed = all(all(row[f].strip() for f in fields) for row in values)
        elif op == "preserve_fields":
            passed = all(len(ids) == 1 and all(row[f] == source[ids[0]][f] for f in fields) for row, ids in zip(values, target["lineage"]))
        elif op in {"scale", "strip_prefix"}:
            passed = True
            for row, ids in zip(values, target["lineage"]):
                if len(ids) != 1 or ids[0] not in source:
                    passed = False
                    break
                raw, actual = source[ids[0]][rule["source_field"]], row[rule["field"]]
                if op == "strip_prefix":
                    passed = passed and actual == raw.removeprefix(rule["prefix"])
                else:
                    try:
                        if not raw.strip():
                            passed = passed and actual == ""
                        else:
                            raw_number, factor, actual_number = Decimal(raw), Decimal(rule["factor"]), Decimal(actual)
                            with localcontext() as context:
                                context.prec = len(raw_number.as_tuple().digits) + len(factor.as_tuple().digits) + 8
                                passed = passed and actual_number.is_finite() and raw_number.is_finite() and factor.is_finite() and actual_number == raw_number * factor
                    except InvalidOperation:
                        passed = False
        else:
            raise ValueError("不支持的项目检查操作：" + str(op))
        check("rule_" + str(index), passed, basis)
    pending = plan.get("pending_checks", [])
    return {"schema": VERSION, "status": "failed" if any(c["status"] == "failed" for c in checks) else "partial" if pending else "passed",
            "checks": checks, "pending_checks": pending, "researcher_acceptance": "not_assessed",
            "runtime": runtime(), "frozen_runtime": plan.get("runtime", "unknown"),
            "input_rows": len(source), "output_rows": sum(len(t["rows"]) for t in output_tables),
            "explained_exclusions": len(explained), "output_files": [{"name": t["file"], "sha256": sha(safe_file(output, t["file"]))} for t in output_tables]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["profile", "run", "verify"])
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--selection", help="JSON input descriptors, only for profile")
    parser.add_argument("--bundle", default=".")
    parser.add_argument("--output-dir", default="replay-output")
    args = parser.parse_args()
    if args.command == "profile":
        result = profile(args.input_dir, json.loads(args.selection))
    elif args.command == "run":
        execution = execute(args.bundle, args.input_dir, args.output_dir)
        if execution.get("questions"):
            print(dumps({"status": "needs_attention", **execution}), end="")
            return 3
        result = verify(args.bundle, args.input_dir, args.output_dir)
        baseline = Path(args.bundle) / "results" / "verification.json"
        if baseline.is_file():
            previous = json.loads(baseline.read_text(encoding="utf-8"))
            equal = previous.get("output_files") == result["output_files"]
            result["checks"].append({"id": "reproduction_output_hashes", "status": "passed" if equal else "failed", "detail": "与复现包原运行的数据文件逐字节比对"})
            if not equal:
                result["status"] = "failed"
        write_json(Path(args.output_dir) / "verification.json", result)
        if result["status"] == "failed":
            print(dumps(result))
            return 2
    else:
        result = verify(args.bundle, args.input_dir, args.output_dir)
    print(dumps(result), end="")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(1)
