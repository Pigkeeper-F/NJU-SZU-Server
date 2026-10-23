"""Literature-mining export stage: 文献表 + 参数明细表 as CSV (and Excel).

Frozen contract: ``docs/literature-mining-stage2-design.md`` §2.4.6 / §3.8 / §6.3 / §6.4.

* inputs  = the whole job chain (candidates / dedupe / fulltext / extraction) plus
  ``validation/validation.json`` and the effective confirmed field definitions;
* outputs = ``export/literature.csv`` + ``literature.xlsx`` (sheet 文献表) and
  ``export/parameters.csv`` + ``parameters.xlsx`` (sheet 参数明细) plus
  ``export_report.md``;
* blocked rows never reach the main tables (they are listed in the report),
  ``pending_review`` rows are written with ``validation_status=pending_review``;
* 版本不覆盖: any existing ``export/`` payload is archived to
  ``export/history/<yyyyMMdd-HHmmss>/`` before the new version is written.
"""
from __future__ import annotations

import argparse
import csv
import importlib
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
TOOLS = Path(__file__).resolve().parent
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from literature_extract import (  # noqa: E402 - same-tool module, path prepared above
    StageError,
    load_config,
    load_effective_schema,
)


def _load_literature_common():
    """引用 I1 的 tools/literature-common.py（设计 §9.1「引用」）。"""
    for name in ("literature_common", "literature-common"):
        try:
            module = importlib.import_module(name)
        except Exception:
            continue
        if callable(getattr(module, "emit", None)):
            return module
    return None


_COMMON = _load_literature_common()


def _shared(name: str, fallback):
    function = getattr(_COMMON, name, None) if _COMMON is not None else None
    return function if callable(function) else fallback


def _fallback_emit(kind: str, **values: object) -> None:
    print(json.dumps({"type": kind, **values}, ensure_ascii=False), flush=True)


def _fallback_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _fallback_atomic_write_text(path: Path, text: str) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="") as stream:
            stream.write(text)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _fallback_atomic_write_json(path: Path, obj: object, *, indent: int = 2) -> None:
    _fallback_atomic_write_text(Path(path), json.dumps(obj, ensure_ascii=False, indent=indent) + "\n")


def _fallback_safe_csv_cell(value: object) -> str:
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _fallback_safe_url(value: object) -> str:
    if not isinstance(value, str):
        return ""
    parsed = urlparse(value.strip())
    return value.strip() if parsed.scheme in {"http", "https"} and parsed.netloc else ""


emit = _shared("emit", _fallback_emit)
now_iso = _shared("now_iso", _fallback_now_iso)
atomic_write_json = _shared("atomic_write_json", _fallback_atomic_write_json)
atomic_write_text = _shared("atomic_write_text", _fallback_atomic_write_text)
safe_csv_cell = _shared("safe_csv_cell", _fallback_safe_csv_cell)
safe_url = _shared("safe_url", _fallback_safe_url)

LITERATURE_COLUMNS = [
    "paper_id", "title", "authors", "year", "doi", "source", "source_url", "fulltext_url",
    "fulltext_status", "evidence_level", "retrieved_at", "screening_status", "screening_reason",
    "review_status", "notes",
]
PARAMETER_COLUMNS = [
    "row_id", "paper_id", "field", "value_raw", "unit_raw", "value_normalized", "unit_normalized",
    "conversion", "evidence_where", "evidence_quote", "provenance", "conditions_json",
    "validation_status", "validation_notes",
]
LITERATURE_SHEET = "文献表"
PARAMETER_SHEET = "参数明细"
FULLTEXT_READY = {"fetched_pdf", "fetched_html"}
EXPORT_FILES = [
    "export/literature.csv", "export/literature.xlsx",
    "export/parameters.csv", "export/parameters.xlsx",
    "export/export_report.md",
]


class ExportError(StageError):
    """Export-stage failure (alias of StageError for readability at call sites)."""


# --------------------------------------------------------------------------- inputs

def _read_json(path: Path) -> object | None:
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None


def load_inputs(job_dir: Path, *, require_validation: bool = True) -> dict:
    """Collect every job artifact the two tables are built from."""
    job_dir = Path(job_dir)
    inputs: dict = {
        "job_dir": job_dir,
        "config": _read_json(job_dir / "run_config.json") or {},
        "candidate_report": _read_json(job_dir / "candidates.json"),
        "candidates": [],
        "groups": [],
        "fulltext": None,
        "bundles": {},
        "extractions": [],
        "validation": None,
        "schema": None,
    }
    report = inputs["candidate_report"]
    if isinstance(report, dict) and isinstance(report.get("records"), list):
        inputs["candidates"] = [record for record in report["records"] if isinstance(record, dict)]
    dedupe = _read_json(job_dir / "dedupe" / "dedupe.json")
    if isinstance(dedupe, dict):
        inputs["groups"] = [group for group in (dedupe.get("groups") or []) if isinstance(group, dict)]
    fulltext = _read_json(job_dir / "fulltext" / "status.json")
    if isinstance(fulltext, dict):
        inputs["fulltext"] = fulltext
    inputs_dir = job_dir / "extraction" / "inputs"
    if inputs_dir.is_dir():
        for path in sorted(inputs_dir.glob("*.json")):
            bundle = _read_json(path)
            if isinstance(bundle, dict):
                inputs["bundles"][path.stem] = bundle
    extraction_root = job_dir / "extraction"
    if extraction_root.is_dir():
        for path in sorted(extraction_root.glob("*.json")):
            if path.name == "check_report.json":
                continue
            document = _read_json(path)
            if isinstance(document, dict):
                inputs["extractions"].append((path.stem, document))
    inputs["validation"] = _read_json(job_dir / "validation" / "validation.json")
    if require_validation and not isinstance(inputs["validation"], dict):
        raise ExportError("尚未完成校验（缺少 validation/validation.json）")
    inputs["schema"] = load_effective_schema(job_dir)
    return inputs


def _fulltext_items(inputs: dict) -> list[dict]:
    if not isinstance(inputs.get("fulltext"), dict):
        return []
    return [item for item in (inputs["fulltext"].get("items") or []) if isinstance(item, dict)]


def _int_or_blank(value: object) -> object:
    return value if isinstance(value, int) and not isinstance(value, bool) else ""


def literature_entries(inputs: dict) -> list[dict]:
    """Ordered literature entries (one per candidate record) before blocking filters.

    Each entry = ``{"canonical_id", "record", "group", "paper_id"}``.  Kept public so
    ``literature_validate`` applies B7/P1 to exactly the rows this module would export.
    """
    ordered: list[dict] = []
    if inputs.get("groups"):
        for group in inputs["groups"]:
            canonical_id = str(group.get("canonical_id") or "")
            members = [record for record in (group.get("records") or []) if isinstance(record, dict)]
            if not members:
                members = [{"candidate_id": canonical_id}]
            for record in members:
                ordered.append({"canonical_id": canonical_id, "record": record, "group": group,
                                "paper_id": str(record.get("candidate_id") or "").strip()})
    else:
        for record in inputs.get("candidates") or []:
            if isinstance(record, dict):
                ordered.append({"canonical_id": str(record.get("candidate_id") or ""), "record": record, "group": None,
                                "paper_id": str(record.get("candidate_id") or "").strip()})
    return ordered


def literature_rows(inputs: dict) -> list[dict]:
    """One row per candidate record (paper), duplicates keep their screening record (§3.8)."""
    ordered = literature_entries(inputs)

    fulltext_by_id = {str(item.get("paper_id") or ""): item for item in _fulltext_items(inputs)}
    bundles = inputs.get("bundles") or {}
    extraction_by_id = {paper_id: document for paper_id, document in inputs.get("extractions") or []}
    validation = inputs.get("validation") if isinstance(inputs.get("validation"), dict) else None
    blocked_ids = set()
    for issue in (validation or {}).get("blocking") or []:
        if isinstance(issue, dict) and issue.get("entity") == "literature":
            blocked_ids.add(str(issue.get("paper_id") or ""))

    rows: list[dict] = []
    seen: set[str] = set()
    for entry in ordered:
        record = entry["record"]
        group = entry["group"] or {}
        canonical_id = entry["canonical_id"]
        paper_id = str(record.get("candidate_id") or "").strip()
        if not paper_id or paper_id in seen or paper_id in blocked_ids:
            continue  # B7: 缺少/重复 paper_id 的行阻断，不写入主表
        seen.add(paper_id)
        is_canonical = paper_id == canonical_id
        item = fulltext_by_id.get(paper_id) if is_canonical else None
        notes: list[str] = []
        member_ids = [str(value) for value in (group.get("candidate_ids") or []) if str(value)]
        if not is_canonical:
            notes.append(f"同组重复/疑似重复：{canonical_id}")
        elif len(member_ids) > 1:
            notes.append(f"归并候选：{'、'.join(member_ids)}")
        if group.get("needs_review") and is_canonical:
            notes.append("疑似重复组（待人工确认，不静默合并）")
        if is_canonical:
            bundle = bundles.get(paper_id)
            if isinstance(bundle, dict) and isinstance(bundle.get("degraded"), dict):
                notes.append(f"全文文本待补充：{bundle['degraded'].get('reason') or ''}")
            document = extraction_by_id.get(paper_id)
            if document is None:
                if isinstance(bundle, dict):
                    notes.append("待补充：尚未完成提取")
                elif item is not None and str(item.get("status") or "") in FULLTEXT_READY:
                    notes.append("待补充：未生成提取输入包")
            else:
                status = str(document.get("status") or "")
                if status == "skipped":
                    notes.append(f"提取跳过：{document.get('skip_reason') or ''}")
                elif status == "partial":
                    notes.append("提取状态：partial（部分字段缺失）")
            if item is not None:
                status = str(item.get("status") or "")
                if status in {"unavailable", "needs_attention"}:
                    notes.append(f"全文待补充：{item.get('reason') or status}")
                elif status == "partial":
                    missing = (item.get("completeness") or {}).get("missing_sections") or []
                    notes.append("正文结构不完整：" + ("、".join(missing) if missing else (item.get("reason") or "缺少核心章节")))
                elif status == "abstract_only":
                    notes.append("仅有摘要：未取得全文正文")
                elif status == "pending":
                    notes.append("全文获取未完成")
                elif not notes and status not in FULLTEXT_READY:
                    notes.append(f"全文状态：{status}")
                materials_state = item.get("materials_completeness") or {}
                if status in FULLTEXT_READY | {"partial"} and materials_state:
                    if materials_state.get("status") == "partial":
                        notes.append("补充材料未下载完整：" + str(materials_state.get("pending") or 0) + " 项待补")
                    elif materials_state.get("status") == "unverified":
                        notes.append("补充材料尚未核查")
        screening_status = str(record.get("screening_status") or ("pending_review" if is_canonical else "auto_excluded_duplicate"))
        screening_reason = str(record.get("screening_reason") or "")
        if not screening_reason:
            if not is_canonical:
                screening_reason = f"duplicate of {canonical_id}"
            elif group.get("needs_review"):
                screening_reason = f"probable duplicate group {canonical_id}（待人工确认）"
        fulltext_url = ""
        if item is not None and str(item.get("status") or "") in FULLTEXT_READY:
            # Candidate links are only unverified leads. A failed/stale URL
            # must not appear in the exported column for acquired full text.
            fulltext_url = safe_url(item.get("fulltext_url"))
        evidence_level = str(record.get("evidence_level") or "")
        if item is not None and str(item.get("evidence_level") or ""):
            evidence_level = str(item["evidence_level"])
        rows.append({
            "paper_id": paper_id,
            "title": str(record.get("title") or group.get("title") or "")[:500],
            "authors": str(record.get("authors") or "")[:500],
            "year": _int_or_blank(record.get("year")) or (_int_or_blank(group.get("year")) if is_canonical else ""),
            "doi": str(record.get("doi") or (group.get("doi") if is_canonical else "") or ""),
            "source": str(record.get("source") or (group.get("sources") or [""])[0] or ""),
            "source_url": safe_url(record.get("result_url")),
            "fulltext_url": fulltext_url,
            "fulltext_status": str(item.get("status") or "") if item is not None else "",
            "evidence_level": evidence_level,
            "retrieved_at": str(record.get("retrieved_at") or (item.get("retrieved_at") if item else "") or ""),
            "screening_status": screening_status,
            "screening_reason": screening_reason,
            "review_status": str(record.get("review_status") or "unreviewed"),
            "notes": "；".join(note for note in notes if note)[:1000],
        })
    return rows


def _row_identity(paper_id: str, row_id: str, index: int) -> tuple[tuple[str, str], tuple[str, int]]:
    return (paper_id, row_id), (paper_id, index)


def _validation_index(inputs: dict) -> tuple[set, set, dict]:
    blocked_ids: set = set()
    blocked_index: set = set()
    pending_notes: dict = {}
    validation = inputs.get("validation") if isinstance(inputs.get("validation"), dict) else None
    for issue in (validation or {}).get("blocking") or []:
        if not isinstance(issue, dict) or issue.get("entity") != "parameter":
            continue
        paper_id = str(issue.get("paper_id") or "")
        row_id = str(issue.get("row_id") or "")
        if row_id:
            blocked_ids.add((paper_id, row_id))
        blocked_index.add((paper_id, int(issue.get("index", -1))))
    for issue in (validation or {}).get("pending") or []:
        if not isinstance(issue, dict) or issue.get("entity") != "parameter":
            continue
        paper_id = str(issue.get("paper_id") or "")
        row_id = str(issue.get("row_id") or "")
        message = str(issue.get("message") or "")
        key = (paper_id, row_id) if row_id else (paper_id, int(issue.get("index", -1)))
        pending_notes[key] = (pending_notes.get(key) or "") + ("；" if pending_notes.get(key) else "") + message
    return blocked_ids, blocked_index, pending_notes


def _cell(value: object) -> object:
    """CSV/Excel 同形的单元格值：None→""、对象→紧凑 JSON、数字/字符串原样。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    return str(value)


def parameter_rows(inputs: dict) -> list[dict]:
    """One row per reported value; blocked rows are dropped (they live in the report)."""
    blocked_ids, blocked_index, pending_notes = _validation_index(inputs)
    rows: list[dict] = []
    for paper_id, document in inputs.get("extractions") or []:
        if str(document.get("status") or "") == "skipped":
            continue
        for index, row in enumerate(document.get("rows") or []):
            if not isinstance(row, dict):
                continue
            row_id = str(row.get("row_id") or "")
            if (paper_id, row_id) in blocked_ids or (paper_id, index) in blocked_index:
                continue
            evidence = row.get("evidence") if isinstance(row.get("evidence"), dict) else {}
            conditions = row.get("conditions") if isinstance(row.get("conditions"), dict) else {}
            key = (paper_id, row_id) if row_id else (paper_id, index)
            note = pending_notes.get(key, "")
            rows.append({
                "row_id": row_id,
                "paper_id": paper_id,
                "field": str(row.get("field") or ""),
                "value_raw": _cell(row.get("value_raw")),
                "unit_raw": _cell(row.get("unit_raw")),
                "value_normalized": _cell(row.get("value_normalized")),
                "unit_normalized": _cell(row.get("unit_normalized")),
                "conversion": _cell(row.get("conversion")),
                "evidence_where": str(evidence.get("where") or ""),
                "evidence_quote": str(evidence.get("quote") or ""),
                "provenance": str(row.get("provenance") or ""),
                "conditions_json": json.dumps(conditions, ensure_ascii=False, separators=(",", ":"), sort_keys=True),
                "validation_status": "pending_review" if note else "ok",
                "validation_notes": note,
            })
    return rows


# --------------------------------------------------------------------------- writers

def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    """utf-8-sig + every cell through ``safe_csv_cell`` (§3.8 形式冻结)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = _temp_path(path)
    try:
        with os.fdopen(handle, "w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(columns))
            writer.writeheader()
            for row in rows:
                writer.writerow({name: safe_csv_cell(_cell(row.get(name))) for name in columns})
        os.replace(temporary, path)
    finally:
        _cleanup(temporary)


def _temp_path(path: Path) -> tuple[int, str]:
    return tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")


def _cleanup(temporary: str) -> None:
    if os.path.exists(temporary):
        os.unlink(temporary)


def normalize_header(value: object) -> str:
    text = str(value or "").strip()
    return re.sub(r"[\s_\-（）()]+", "", text).casefold()


def map_sheet(worksheet, columns: list[str]) -> dict:
    """Map template columns to table columns (单行表头按列名 / 双行表头走 group.field 机制)."""
    merged_row1 = [entry for entry in getattr(worksheet, "merged_cells", None).ranges if entry.min_row == 1 and entry.max_row == 1] if getattr(worksheet, "merged_cells", None) else []
    header_rows = 2 if merged_row1 else 1
    by_normalized = {normalize_header(column): column for column in columns}
    mapping: dict[int, str] = {}
    unmatched: list[str] = []
    row1_group: dict[int, str] = {}
    for entry in merged_row1:
        top_left = worksheet.cell(1, entry.min_col).value
        for index in range(entry.min_col, entry.max_col + 1):
            row1_group[index] = str(top_left or "")
    for column_index in range(1, (worksheet.max_column or 0) + 1):
        row1 = worksheet.cell(1, column_index).value
        row2 = worksheet.cell(2, column_index).value if header_rows == 2 else None
        candidates: list[str] = []
        header_label = ""
        if header_rows == 2:
            group = row1_group.get(column_index) or (str(row1).strip() if row1 not in (None, "") else "")
            if row2 not in (None, ""):
                if group:
                    candidates.append(f"{group}.{row2}")
                candidates.append(str(row2))
                header_label = f"{group}.{row2}" if group else str(row2)
            elif group:
                candidates.append(group)
                header_label = group
        else:
            if row1 is not None:
                candidates.append(str(row1))
            header_label = str(row1 or "")
        if not candidates:
            continue
        matched = None
        for candidate in candidates:
            found = by_normalized.get(normalize_header(candidate))
            if found:
                matched = found
                break
        if matched:
            mapping[column_index] = matched
        elif header_label.strip("."):
            unmatched.append(header_label)
    used = set(mapping.values())
    return {
        "header_rows": header_rows,
        "mapping": {int(index): name for index, name in sorted(mapping.items())},
        "unmapped_columns": [column for column in columns if column not in used],
        "unmapped_template": unmatched,
    }


def _import_openpyxl():
    try:
        import openpyxl  # noqa: PLC0415 - optional dependency, imported lazily
    except Exception as error:
        raise ExportError(f"导出 Excel 需要 openpyxl（当前不可用：{error}）") from None
    return openpyxl


def write_xlsx(path: Path, rows: list[dict], columns: list[str], sheet: str, template: Path | None) -> None:
    """Excel twin of ``write_csv``; ``template`` (optional) supplies header + mapping."""
    openpyxl = _import_openpyxl()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if template is not None:
        workbook = openpyxl.load_workbook(str(template))
        if sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]
        layout = map_sheet(worksheet, columns)
        header_rows = int(layout["header_rows"])
        if worksheet.max_row > header_rows:
            worksheet.delete_rows(header_rows + 1, worksheet.max_row - header_rows)
        for offset, row in enumerate(rows):
            target = header_rows + 1 + offset
            for column_index, name in layout["mapping"].items():
                worksheet.cell(row=target, column=int(column_index), value=_cell(row.get(name)))
    else:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = sheet
        worksheet.append(list(columns))
        for row in rows:
            worksheet.append([_cell(row.get(name)) for name in columns])
    handle, temporary = _temp_path(path)
    os.close(handle)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        _cleanup(temporary)


def archive_previous(export_dir: Path) -> str | None:
    """Move the previous export payload into ``export/history/<yyyyMMdd-HHmmss>/`` (§6.3)."""
    export_dir = Path(export_dir)
    if not export_dir.is_dir():
        return None
    entries = [entry for entry in export_dir.iterdir() if entry.name != "history"]
    if not entries:
        return None
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    history = export_dir / "history" / stamp
    suffix = 0
    while history.exists():
        suffix += 1
        history = export_dir / "history" / f"{stamp}-{suffix}"
    history.mkdir(parents=True)
    for entry in entries:
        shutil.move(str(entry), str(history / entry.name))
    return f"history/{history.name}"


def resolve_template(value: object) -> Path | None:
    if value in (None, "", False):
        return None
    if not isinstance(value, str):
        raise ExportError("运行配置无效：templates.* 必须是相对项目根的路径字符串")
    candidate = Path(value)
    if not candidate.is_absolute():
        candidate = ROOT / candidate
    if not candidate.is_file():
        raise ExportError(f"用户模板不可用：{value}")
    return candidate


# --------------------------------------------------------------------------- report

def _fill_rates(rows: list[dict], columns: list[str]) -> dict[str, float]:
    if not rows:
        return {column: 0.0 for column in columns}
    return {column: sum(1 for row in rows if str(_cell(row.get(column))).strip() != "") / len(rows) for column in columns}


def build_report(inputs: dict, literature: list[dict], parameters: list[dict], archived: str | None,
                 template_layouts: dict | None = None) -> str:
    validation = inputs.get("validation") if isinstance(inputs.get("validation"), dict) else {}
    counts = validation.get("counts") if isinstance(validation.get("counts"), dict) else {}
    schema = inputs.get("schema") or {}
    lines: list[str] = []
    lines.append("# 导出报告（文献表 + 参数明细表）")
    lines.append("")
    lines.append(f"- 生成时间：{now_iso()}")
    lines.append(f"- 字段定义来源：{schema.get('source') or '（未确认字段定义）'}")
    lines.append(f"- 字段定义指纹 schema_sha256：{(schema.get('sha256') or '')[:16]}")
    lines.append("- 两表依据同一份已确认字段定义生成；被排除文献保留筛选记录，阻断行只在下方错误报告列出。")
    lines.append("")
    lines.append("## 计数摘要")
    lines.append("")
    lines.append(f"- 文献表行数：{len(literature)}")
    lines.append(f"- 参数明细表行数：{len(parameters)}")
    lines.append(f"- 参数行校验：共 {counts.get('parameters_total', '')}，通过 {counts.get('ok', '')}，阻断 {counts.get('blocking', '')}，待审核 {counts.get('pending', '')}")
    lines.append(f"- 文献行阻断：{counts.get('literature_blocking', 0)}，文献行待审核：{counts.get('literature_pending', 0)}")
    status_histogram: dict[str, int] = {}
    for row in literature:
        key = str(row.get("fulltext_status") or "未获取")
        status_histogram[key] = status_histogram.get(key, 0) + 1
    if status_histogram:
        lines.append("- 文献全文状态：" + "，".join(f"{key} {value}" for key, value in sorted(status_histogram.items())))
    pending_rows = sum(1 for row in parameters if row.get("validation_status") == "pending_review")
    lines.append(f"- 参数明细表待审核行：{pending_rows}")
    if archived:
        lines.append(f"- 旧版本已归档：export/{archived}")
    lines.append("")
    lines.append("## 列填充率")
    lines.append("")
    for title, rows, columns in (("文献表", literature, LITERATURE_COLUMNS), ("参数明细表", parameters, PARAMETER_COLUMNS)):
        rates = _fill_rates(rows, columns)
        low = [f"{column} {rate:.0%}" for column, rate in rates.items() if rate < 0.1]
        lines.append(f"- {title}：低于 10% 填充的列 {'、'.join(low) if low else '（无）'}")
    lines.append("")
    lines.append("## 错误报告：阻断行（未写入主表）")
    lines.append("")
    blocking = [issue for issue in (validation.get("blocking") or []) if isinstance(issue, dict)]
    if not blocking:
        lines.append("（无阻断行）")
    else:
        lines.append("| 类型 | 记录 | 论文 | 规则 | 问题 |")
        lines.append("|---|---|---|---|---|")
        for issue in blocking:
            record = str(issue.get("row_id") or "-")
            lines.append("| {entity} | {record} | {paper} | {rules} | {message} |".format(
                entity=str(issue.get("entity") or ""),
                record=record.replace("|", "｜"),
                paper=str(issue.get("paper_id") or "").replace("|", "｜"),
                rules="、".join(issue.get("rules") or []),
                message=str(issue.get("message") or "").replace("|", "｜"),
            ))
    lines.append("")
    lines.append("## 待审核行（已写入主表，另见 validation.md）")
    lines.append("")
    pending = [issue for issue in (validation.get("pending") or []) if isinstance(issue, dict)]
    if not pending:
        lines.append("（无待审核行）")
    else:
        lines.append("| 类型 | 记录 | 论文 | 规则 | 问题 |")
        lines.append("|---|---|---|---|---|")
        for issue in pending:
            lines.append("| {entity} | {record} | {paper} | {rules} | {message} |".format(
                entity=str(issue.get("entity") or ""),
                record=str(issue.get("row_id") or "-").replace("|", "｜"),
                paper=str(issue.get("paper_id") or "").replace("|", "｜"),
                rules="、".join(issue.get("rules") or []),
                message=str(issue.get("message") or "").replace("|", "｜"),
            ))
    lines.append("")
    lines.append("## 待补充 / 降级清单")
    lines.append("")
    supplements = [(row["paper_id"], row["notes"]) for row in literature if "待补充" in str(row.get("notes") or "") or "跳过" in str(row.get("notes") or "")]
    if not supplements:
        lines.append("（无）")
    else:
        for paper_id, note in supplements:
            lines.append(f"- {paper_id}：{note}")
    lines.append("")
    if template_layouts:
        lines.append("## 用户模板映射（供研究者确认）")
        lines.append("")
        for table_name, layout in template_layouts.items():
            lines.append(f"### {table_name}")
            lines.append("")
            lines.append(f"- 表头行数：{layout.get('header_rows')}；映射列数：{len(layout.get('mapping') or {})}")
            lines.append("- 映射：" + ("；".join(f"列{index}→{name}" for index, name in (layout.get("mapping") or {}).items()) or "（无）"))
            if layout.get("unmapped_columns"):
                lines.append("- 模板未覆盖的表列：" + "、".join(layout["unmapped_columns"]))
            if layout.get("unmapped_template"):
                lines.append("- 未识别的模板表头：" + "、".join(layout["unmapped_template"]))
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


# --------------------------------------------------------------------------- run

def run(config: dict, job_dir: Path) -> dict:
    job_dir = Path(job_dir)
    inputs = load_inputs(job_dir)
    literature = literature_rows(inputs)
    parameters = parameter_rows(inputs)
    export_dir = job_dir / "export"
    export_dir.mkdir(parents=True, exist_ok=True)
    archived = archive_previous(export_dir)

    templates = config.get("templates") if isinstance(config, dict) else None
    templates = templates if isinstance(templates, dict) else {}
    literature_template = resolve_template(templates.get("literature"))
    parameters_template = resolve_template(templates.get("parameters"))
    template_layouts: dict[str, dict] = {}
    for table_name, template, sheet in (("文献表", literature_template, LITERATURE_SHEET), ("参数明细表", parameters_template, PARAMETER_SHEET)):
        if template is None:
            continue
        openpyxl = _import_openpyxl()
        workbook = openpyxl.load_workbook(str(template))
        worksheet = workbook[sheet] if sheet in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        template_layouts[table_name] = map_sheet(worksheet, LITERATURE_COLUMNS if table_name == "文献表" else PARAMETER_COLUMNS)

    write_csv(export_dir / "literature.csv", literature, LITERATURE_COLUMNS)
    write_xlsx(export_dir / "literature.xlsx", literature, LITERATURE_COLUMNS, LITERATURE_SHEET, literature_template)
    write_csv(export_dir / "parameters.csv", parameters, PARAMETER_COLUMNS)
    write_xlsx(export_dir / "parameters.xlsx", parameters, PARAMETER_COLUMNS, PARAMETER_SHEET, parameters_template)

    validation = inputs.get("validation") if isinstance(inputs.get("validation"), dict) else {}
    validation_counts = validation.get("counts") if isinstance(validation.get("counts"), dict) else {}
    counts = {
        "literature": len(literature),
        "parameters": len(parameters),
        "pending_review_parameters": sum(1 for row in parameters if row.get("validation_status") == "pending_review"),
        "blocking": validation_counts.get("blocking", 0),
        "pending": validation_counts.get("pending", 0),
    }
    atomic_write_text(export_dir / "export_report.md", build_report(inputs, literature, parameters, archived, template_layouts))
    return {"count": len(parameters), "counts": counts, "files": list(EXPORT_FILES), "archived": archived, "literature": len(literature)}


# --------------------------------------------------------------------------- CLI

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="文献挖掘 export 阶段：文献表 + 参数明细表（CSV/Excel，版本不覆盖）")
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--resume", action="store_true", help="与普通运行等价（export 无条件全量重算，仅作 CLI 对称性）")
    args = parser.parse_args(argv)
    try:
        config = load_config(args.config)
        emit("stage", stage="export", message="按已确认字段定义导出文献表与参数明细表")
        result = run(config, args.output_dir)
        message = f"已导出 {result['literature']} 篇文献、{result['count']} 个参数行"
        if result["archived"]:
            message += f"；旧版本已归档到 export/{result['archived']}"
        emit("result", stage="export", message=message, count=result["count"], counts=result["counts"], files=result["files"])
        return 0
    except StageError as error:
        emit("error", message=str(error))
        return 1
    except Exception as error:  # noqa: BLE001 - stage boundary: any failure becomes one error event
        emit("error", message=f"导出阶段失败：{error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
