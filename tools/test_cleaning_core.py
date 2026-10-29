"""Offline cleaning contracts; disposable fixtures, never rewrite source samples."""
import csv
import importlib.util
import json
import shutil
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch
from pathlib import Path

import cleaning_core as core

TOOLS = Path(__file__).resolve().parent
IDENTITY = '''def clean(tables, config):
    return {"tables": [{"name": t["id"], "columns": t["columns"], "rows": t["rows"],
                        "lineage": [[s] for s in t["source_ids"]]} for t in tables], "dispositions": {}}
'''


class CleaningCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs, self.bundle = self.root / "inputs", self.root / "package"
        self.inputs.mkdir()
        self.bundle.mkdir()

    def fixture(self, text="code,value\n004,-11.000000\n005,0.000000\n", name="data.csv"):
        (self.inputs / name).write_text(text, encoding="utf-8-sig")
        return [{"name": name}]

    def freeze(self, selection, code=IDENTITY, checks=None):
        profile = core.profile(self.inputs, selection)
        shutil.copyfile(TOOLS / "cleaning_core.py", self.bundle / "cleaning_core.py")
        (self.bundle / "clean.py").write_text(code, encoding="utf-8")
        plan = {"schema": 1, "inputs": profile["inputs"], "rules": {"text": "保留原值；不得丢行；两行", "version": 1},
                "script_sha256": core.sha(self.bundle / "clean.py"), "engine_sha256": core.sha(self.bundle / "cleaning_core.py"),
                "checks": checks or [], "pending_checks": []}
        core.write_json(self.bundle / "cleaning-plan.json", plan)
        return plan

    def test_fbsh_profile_matches_readonly_baseline(self):
        sample = TOOLS.parent / "examples/1992-cleaning-trial/FBSH_1992_Afghanistan_练习样本.csv"
        shutil.copyfile(sample, self.inputs / "sample.csv")
        before = core.sha(sample)
        p = core.profile(self.inputs, [{"name": "sample.csv"}])
        t = p["tables"][0]
        fields = {f["name"]: f for f in t["fields"]}
        self.assertFalse(p["blocked"])
        self.assertEqual((p["row_count"], len(t["columns"])), (80, 13))
        self.assertEqual((fields["Value"]["zero_count"], fields["Value"]["negative_count"]), (20, 2))
        self.assertEqual(fields["Flag"]["values"], {"E": 30, "I": 50})
        self.assertEqual(fields["Item Code"]["distinct"], 23)
        self.assertEqual(core.sha(sample), before)

    def test_csv_quotes_multiline_tsv_and_blank_are_not_silently_dropped(self):
        select = self.fixture('code,note\n004,"a,b\nsecond line"\n\n')
        p = core.profile(self.inputs, select)
        self.assertTrue(p["blocked"])
        self.assertEqual(p["row_count"], 2)
        self.assertEqual(p["tables"][0]["sample"][0][1], "a,b\nsecond line")
        tsv = self.fixture("code\tvalue\n004\t0\n", "data.tsv")
        self.assertFalse(core.profile(self.inputs, tsv)["blocked"])

    def test_reproduction_outputs_are_byte_identical_and_keep_codes_and_decimal(self):
        select = self.fixture()
        self.freeze(select, checks=[{"op": "preserve_fields", "table": 1, "fields": ["code", "value"], "basis": "保留原值"}])
        for name in ("one", "two"):
            out = self.root / name
            core.execute(self.bundle, self.inputs, out)
            report = core.verify(self.bundle, self.inputs, out)
            self.assertEqual(report["status"], "passed")
            self.assertEqual(report["researcher_acceptance"], "not_assessed")
            self.assertEqual(report["runtime"]["python"], core.runtime()["python"])
            self.assertEqual(report["frozen_runtime"], "unknown")
            with (out / "table-01.csv").open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(rows[0], {"code": "004", "value": "-11.000000"})
        for name in ("table-01.csv", "cleaning-result.json"):
            self.assertEqual(core.sha(self.root / "one" / name), core.sha(self.root / "two" / name))

    def test_fbsh_project_script_uses_generic_runner_and_preserves_raw_fields(self):
        sample = TOOLS.parent / "examples/1992-cleaning-trial/FBSH_1992_Afghanistan_练习样本.csv"
        shutil.copyfile(sample, self.inputs / "sample.csv")
        code = '''from decimal import Decimal
def clean(tables, config):
    t=tables[0]
    rows=[dict(r,Country_M49=r["Area Code (M49)"].removeprefix("'"),
               Value_t=str(Decimal(r["Value"])*1000)) for r in t["rows"]]
    return {"tables":[{"columns":t["columns"]+["Country_M49","Value_t"],"rows":rows,
                       "lineage":[[s] for s in t["source_ids"]]}],"dispositions":{}}
'''
        raw_columns = core.profile(self.inputs, [{"name": "sample.csv"}])["tables"][0]["columns"]
        self.freeze([{"name": "sample.csv"}], code=code, checks=[
            {"op": "preserve_fields", "table": 1, "fields": raw_columns, "basis": "保留原值"},
            {"op": "scale", "table": 1, "field": "Value_t", "source_field": "Value", "factor": "1000", "basis": "保留原值"},
            {"op": "strip_prefix", "table": 1, "field": "Country_M49", "source_field": "Area Code (M49)", "prefix": "'", "basis": "保留原值"}])
        out = self.root / "out"
        core.execute(self.bundle, self.inputs, out)
        report = core.verify(self.bundle, self.inputs, out)
        self.assertEqual((report["status"], report["input_rows"], report["output_rows"]), ("passed", 80, 80))
        data = json.loads((out / "cleaning-result.json").read_text())["tables"][0]["rows"]
        self.assertEqual(data[43]["Value_t"], "-11000.000000")
        self.assertTrue(all(r["Country_M49"] == "004" for r in data))

    def test_independent_scale_check_detects_wrong_conversion(self):
        code = '''def clean(tables, config):
    t=tables[0]
    return {"tables":[{"columns":t["columns"]+["tonnes"],"rows":[dict(r,tonnes="0") for r in t["rows"]],
                       "lineage":[[s] for s in t["source_ids"]]}]}
'''
        self.freeze(self.fixture(), code=code, checks=[{"op": "scale", "table": 1, "field": "tonnes", "source_field": "value", "factor": "1000", "basis": "保留原值"}])
        out = self.root / "out"
        core.execute(self.bundle, self.inputs, out)
        report = core.verify(self.bundle, self.inputs, out)
        self.assertEqual(report["status"], "failed")
        self.assertEqual(report["checks"][-1]["id"], "rule_1")

    def test_independent_verifier_detects_unexplained_loss(self):
        self.freeze(self.fixture(), code='''def clean(tables, config):
    t=tables[0]
    return {"tables":[{"columns":t["columns"],"rows":t["rows"][:1],"lineage":[[t["source_ids"][0]]]}]}
''')
        output = self.root / "output"
        core.execute(self.bundle, self.inputs, output)
        report = core.verify(self.bundle, self.inputs, output)
        self.assertEqual(report["status"], "failed")
        self.assertIn("source_records_accounted", [c["id"] for c in report["checks"] if c["status"] == "failed"])

    def test_independent_verifier_detects_csv_tampering_and_unapproved_check(self):
        self.freeze(self.fixture())
        output = self.root / "output"
        core.execute(self.bundle, self.inputs, output)
        (output / "table-01.csv").write_text("code,value\n004,1000\n", encoding="utf-8")
        self.assertEqual(core.verify(self.bundle, self.inputs, output)["status"], "failed")
        plan = json.loads((self.bundle / "cleaning-plan.json").read_text())
        plan["checks"] = [{"op": "row_count", "table": 1, "expected": 2, "basis": "researcher never approved this"}]
        core.write_json(self.bundle / "cleaning-plan.json", plan)
        with self.assertRaisesRegex(ValueError, "原句"):
            core.verify(self.bundle, self.inputs, output)

    def test_hashes_reject_changed_input_and_script_before_execution(self):
        self.freeze(self.fixture())
        (self.inputs / "data.csv").write_text("changed", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "指纹"):
            core.execute(self.bundle, self.inputs, self.root / "out")
        (self.bundle / "clean.py").write_text("raise RuntimeError('must never run')", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "脚本已变化"):
            core.execute(self.bundle, self.inputs, self.root / "out")

    def test_runtime_scientific_question_pauses_without_fake_outputs(self):
        self.freeze(self.fixture(), code='def clean(tables, config):\n    return {"questions":["请明确新出现的单位如何处理"]}\n')
        output = self.root / "output"
        result = core.execute(self.bundle, self.inputs, output)
        self.assertEqual(result["questions"], ["请明确新出现的单位如何处理"])
        self.assertTrue((output / "needs-attention.json").is_file())
        self.assertFalse((output / "cleaning-result.json").exists())

    def test_paths_and_duplicate_selection_rejected(self):
        self.fixture()
        for value in ([{"name": "../data.csv"}], [{"name": "data.csv"}, {"name": "data.csv"}]):
            with self.assertRaises(ValueError):
                core.profile(self.inputs, value)
        (self.inputs / "linked.csv").symlink_to(self.inputs / "data.csv")
        with self.assertRaises(ValueError):
            core.profile(self.inputs, [{"name": "linked.csv"}])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl not installed")
    def test_excel_sheet_selection_formula_cache_and_merged_headers(self):
        import openpyxl
        book = openpyxl.Workbook()
        good = book.active
        good.title = "Data"
        good.append(["code", "value"])
        good.append(["004", 0])
        bad = book.create_sheet("Needs review")
        bad.append(["code", "value"])
        bad.append(["005", "=1+1"])
        merged = book.create_sheet("Merged")
        merged.append(["heading", "x"])
        merged.merge_cells("A1:B1")
        book.save(self.inputs / "sample.xlsx")
        all_tables = core.profile(self.inputs, [{"name": "sample.xlsx"}])
        self.assertTrue(all_tables["blocked"])
        codes = [issue["code"] for t in all_tables["tables"] for issue in t["issues"]]
        self.assertIn("formula_cache_missing", codes)
        self.assertIn("merged_cells", codes)
        good_only = core.profile(self.inputs, [{"name": "sample.xlsx", "sheets": ["Data"]}])
        self.assertFalse(good_only["blocked"])
        self.assertEqual(good_only["tables"][0]["sample"], [["004", "0"]])
        self.assertEqual(good_only["inputs"][0]["available_sheets"], ["Data", "Needs review", "Merged"])

    @unittest.skipUnless(importlib.util.find_spec("xlrd"), "xlrd not installed")
    def test_xls_error_boolean_and_date_types_are_not_silently_numeric(self):
        # Stub the reader, not the cleaning logic; no additional XLS writer dependency.
        import xlrd
        values = [["code", "value"], ["004", True], ["005", 7], ["006", 44927.0]]
        kinds = [[1, 1], [1, xlrd.XL_CELL_BOOLEAN], [1, xlrd.XL_CELL_ERROR], [1, xlrd.XL_CELL_DATE]]
        sheet = SimpleNamespace(nrows=4, ncols=2, merged_cells=[],
                                cell_value=lambda r, c: values[r][c], cell_type=lambda r, c: kinds[r][c])
        book = SimpleNamespace(sheet_names=lambda: ["Data"], sheet_by_name=lambda _: sheet,
                               datemode=0, release_resources=lambda: None)
        (self.inputs / "sample.xls").write_bytes(b"mock reader fixture")
        with patch.object(xlrd, "open_workbook", return_value=book):
            report = core.profile(self.inputs, [{"name": "sample.xls"}])
        self.assertTrue(report["blocked"])
        table = report["tables"][0]
        self.assertEqual(table["sample"], [["004", "true"], ["005", "#DIV/0!"], ["006", "44927.0"]])
        self.assertEqual(table["metadata"]["date_cells"], ["R4C2"])
        self.assertEqual(table["metadata"]["date_mode"], 0)
        self.assertTrue(any(i["code"] == "excel_error_cells" and i["blocking"] for i in table["issues"]))


if __name__ == "__main__":
    unittest.main()
