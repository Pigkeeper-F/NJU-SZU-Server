"""Read research materials without executing embedded code or accessing the network."""
import csv
import io
import json
import pathlib
import sys
import zipfile
import xml.etree.ElementTree as ET

MAX_TEXT = 120000
MAX_UNZIP = 40 * 1024 * 1024


def text(node):
    return " ".join("".join(node.itertext()).split())


def read_bytes(data, name):
    ext = pathlib.Path(name).suffix.lower()
    segments = []
    warnings = []
    if ext == ".xml":
        root = ET.fromstring(data)
        body = root.find(".//body")
        if body is None:
            raise ValueError("XML 没有正文 body，不能当作全文提取")
        for i, node in enumerate(body.iter("p"), 1):
            if text(node):
                segments.append({"location": f"正文段落:{i}", "text": text(node)})
        for i, table in enumerate(body.iter("table-wrap"), 1):
            table_id = table.get("id") or str(i)
            label = table.find("label")
            caption = table.find("caption")
            segments.append({"location": f"表:{table_id}:说明", "text": " ".join(text(n) for n in (label, caption) if n is not None)})
            for j, row in enumerate(table.iter("tr"), 1):
                cells = [text(cell) for cell in row if cell.tag in ("td", "th")]
                segments.append({"location": f"表:{table_id}:行:{j}", "text": " | ".join(cells)})
        for i, fig in enumerate(body.iter("fig"), 1):
            segments.append({"location": f"图:{fig.get('id') or i}:图注", "text": text(fig)})
        abstract = root.find(".//abstract")
        if abstract is not None:
            segments.insert(0, {"location": "摘要（不能替代正文证据）", "text": text(abstract)})
    elif ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("文件内容不是 PDF，请检查是否下载到了网页")
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages]
        except ImportError:
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    pages = [(page.extract_text() or "") for page in pdf.pages]
            except ImportError:
                import fitz
                doc = fitz.open(stream=data, filetype="pdf")
                pages = [page.get_text() for page in doc]
        for i, page in enumerate(pages, 1):
            segments.append({"location": f"PDF页:{i}", "text": page.strip()})
        if not any(s["text"].strip() for s in segments):
            raise ValueError("PDF 为扫描件或无法提取文字，请提供可读材料；当前不自动 OCR")
        if any(not s["text"].strip() for s in segments):
            warnings.append("PDF 有无文字页；图像及图中数值不能按文字提取，需人工检查")
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in book:
            for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = ["" if v is None else str(v) for v in row]
                segments.append({"location": f"工作表:{sheet.title}:行:{i}", "text": " | ".join(cells)})
        book.close()
    elif ext in (".csv", ".tsv"):
        decoded = data.decode("utf-8-sig")
        for i, row in enumerate(csv.reader(io.StringIO(decoded), delimiter="\t" if ext == ".tsv" else ","), 1):
            segments.append({"location": f"表格行:{i}", "text": " | ".join(row)})
    elif ext in (".txt", ".md"):
        decoded = data.decode("utf-8-sig")
        for i, line in enumerate(decoded.splitlines(), 1):
            if line.strip():
                segments.append({"location": f"文本行:{i}", "text": line.strip()})
    elif ext == ".zip":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(e.file_size for e in entries) > MAX_UNZIP:
                raise ValueError("附件 ZIP 解压内容超过小流程限制，请提供需要处理的单个文件")
            for entry in entries:
                if entry.is_dir():
                    continue
                if pathlib.Path(entry.filename).suffix.lower() not in (".pdf", ".txt", ".md", ".csv", ".tsv", ".xlsx", ".xml"):
                    warnings.append(f"附件 {entry.filename} 尚不支持自动读取（图像需人工读图）")
                    continue
                try:
                    parsed = read_bytes(archive.read(entry), entry.filename)
                    for segment in parsed[0]:
                        segment["location"] = f"附件:{entry.filename}/" + segment["location"]
                    segments.extend(parsed[0])
                    warnings.extend(parsed[1])
                except Exception as error:
                    warnings.append(f"附件 {entry.filename} 无法读取：{error}")
    else:
        raise ValueError("材料格式尚不支持，请提供 PDF、XML、文本或表格")
    if sum(len(s["text"]) for s in segments) > MAX_TEXT:
        raise ValueError("材料文字超过 12 万字符；未截断提取，请提供本次所需章节或表格")
    if not any(s["text"] for s in segments):
        raise ValueError("没有可读取的正文或表格")
    return segments, warnings


if __name__ == "__main__":
    try:
        source = pathlib.Path(sys.argv[1])
        if source.stat().st_size > 20 * 1024 * 1024:
            raise ValueError("材料超过 20MB")
        parts, warnings = read_bytes(source.read_bytes(), source.name)
        print(json.dumps({"segments": parts, "warnings": warnings}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        sys.exit(1)
