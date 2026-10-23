"""Tests for tools/literature_export.py (export 阶段：两表 CSV/Excel + 版本不覆盖).

Run (isolated temp cwd):  python -m unittest -v test_literature_export
Coverage: 文献表/参数明细表字段与 CSV 输出（utf-8-sig + 防公式注入）、xlsx 生成、
阻断行不写入主表、待审核行写入并标注、conditions_json 紧凑同形、
同 job 重导出归档到 export/history/、用户模板映射、缺 validation/缺模板的报错。
"""
from __future__ import annotations

import contextlib
import csv
import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

import literature_export as export  # noqa: E402
import literature_validate as validate  # noqa: E402

from test_literature_validate import (  # noqa: E402
    FIELDS,
    base_docs,
    base_groups,
    base_items,
    base_records,
    copy,
    row,
    write_job,
)

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except Exception:  # pragma: no cover - 环境缺 openpyxl 时按 zip 结构级校验
    HAS_OPENPYXL = False


def run_cli(main, argv: list[str]) -> tuple[int, list[dict]]:
    stream = io.StringIO()
    with contextlib.redirect_stdout(stream):
        code = main(argv)
    return code, [json.loads(line) for line in stream.getvalue().splitlines() if line.strip()]


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="lit-i4-export-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def build_job(self, name="job", *, docs=None, records=None, groups="default", items=None, fields=None):
        if records is None:
            records = base_records()
        if groups == "default":
            groups = base_groups()
        if fields is None:
            fields = FIELDS
        job = write_job(self.root / name, records=records, groups=groups, items=items, docs=docs, fields=fields)
        code, events = run_cli(validate.main, ["--config", str(job / "run_config.json"), "--output-dir", str(job)])
        assert code == 0 and events[-1]["type"] == "result", events
        return job

    def default_job(self, name="job"):
        """含 1 个阻断行（value_raw 空）+ 1 个待审核行（单位空）+ 其余通过。"""
        docs = base_docs()
        docs["GSC-001"] = copy.deepcopy(docs["GSC-001"])
        docs["GSC-001"]["rows"][0]["value_raw"] = ""      # B4 阻断
        docs["GSC-002"] = copy.deepcopy(docs["GSC-002"])
        docs["GSC-002"]["rows"][0]["unit_raw"] = ""        # P2 待审核
        docs["GSC-002"]["rows"][0]["unit_normalized"] = ""
        return self.build_job(name, docs=docs)

    def export(self, job, extra=()):
        argv = ["--config", str(job / "run_config.json"), "--output-dir", str(job), *extra]
        return run_cli(export.main, argv)

    # ------------------------------------------------------------------ tables

    def test_failed_legacy_fulltext_link_is_not_exported_as_verified(self):
        other = "https://pmc.ncbi.nlm.nih.gov/articles/PMC10522860/"
        inputs = {
            "candidates": [{"candidate_id": "PMID-010", "title": "Target paper",
                            "result_url": "https://pubmed.ncbi.nlm.nih.gov/33132728/",
                            "fulltext_links": [{"url": other}]}],
            "fulltext": {"items": [{"paper_id": "PMID-010", "status": "unavailable",
                                    "fulltext_url": other, "reason": "Access failed"}]},
        }
        rows = export.literature_rows(inputs)
        self.assertEqual(rows[0]["fulltext_url"], "")
        self.assertEqual(rows[0]["fulltext_status"], "unavailable")

    def test_export_writes_pair_of_csv_and_xlsx(self):
        job = self.default_job()
        code, events = self.export(job)
        self.assertEqual(code, 0)
        final = events[-1]
        self.assertEqual(final["type"], "result")
        self.assertEqual(final["files"], export.EXPORT_FILES)
        self.assertEqual(final["count"], 2)  # 3 个参数行 - 1 阻断
        self.assertEqual(final["counts"]["blocking"], 1)
        self.assertEqual(final["counts"]["pending"], 1)
        for name in export.EXPORT_FILES:
            self.assertTrue((job / name).is_file(), name)

    def test_literature_csv_fields_and_screening_records(self):
        job = self.default_job(name="litcsv")
        self.export(job)
        raw = (job / "export" / "literature.csv").read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "utf-8-sig BOM 缺失")
        with (job / "export" / "literature.csv").open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            self.assertEqual(reader.fieldnames, export.LITERATURE_COLUMNS)
            rows = {entry["paper_id"]: entry for entry in reader}
        self.assertEqual(set(rows), {"GSC-001", "GSC-B2-003", "GSC-002"})
        duplicate = rows["GSC-B2-003"]
        self.assertEqual(duplicate["screening_status"], "auto_excluded_duplicate")
        self.assertIn("GSC-001", duplicate["screening_reason"])
        self.assertEqual(duplicate["review_status"], "unreviewed")
        canonical = rows["GSC-001"]
        self.assertEqual(canonical["fulltext_status"], "fetched_html")
        self.assertEqual(canonical["evidence_level"], "fulltext_html")
        self.assertEqual(canonical["screening_status"], "pending_review")
        self.assertEqual(canonical["source"], "Google Scholar")
        self.assertEqual(canonical["source_url"], "https://example.org/x")
        self.assertIn("归并候选", canonical["notes"])
        unavailable = rows["GSC-002"]
        self.assertEqual(unavailable["fulltext_status"], "unavailable")
        self.assertIn("全文待补充", unavailable["notes"])

    def test_parameters_csv_excludes_blocked_and_flags_pending(self):
        job = self.default_job(name="paramcsv")
        self.export(job)
        with (job / "export" / "parameters.csv").open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            self.assertEqual(reader.fieldnames, export.PARAMETER_COLUMNS)
            rows = list(reader)
        self.assertEqual(len(rows), 2)
        self.assertNotIn("GSC-001-R001", [entry["row_id"] for entry in rows])
        pending = [entry for entry in rows if entry["validation_status"] == "pending_review"]
        self.assertEqual([entry["row_id"] for entry in pending], ["GSC-002-R001"])
        self.assertIn("单位", pending[0]["validation_notes"])
        ok_row = next(entry for entry in rows if entry["row_id"] == "GSC-001-R002")
        self.assertEqual(ok_row["validation_status"], "ok")
        self.assertEqual(ok_row["validation_notes"], "")
        self.assertEqual(ok_row["evidence_quote"], "continuous maize")
        self.assertEqual(ok_row["conditions_json"], '{"site":"A"}')

    def test_report_lists_blocking_and_counts(self):
        job = self.default_job(name="report")
        self.export(job)
        report = (job / "export" / "export_report.md").read_text(encoding="utf-8")
        self.assertIn("## 计数摘要", report)
        self.assertIn("## 错误报告：阻断行（未写入主表）", report)
        self.assertIn("GSC-001-R001", report)
        self.assertIn("missing_value", report)
        self.assertIn("## 待审核行", report)
        self.assertIn("## 待补充 / 降级清单", report)
        self.assertIn("两表依据同一份已确认字段定义生成", report)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl 未安装")
    def test_xlsx_sheets_and_cells(self):
        import openpyxl  # noqa: PLC0415
        job = self.default_job(name="xlsx")
        self.export(job)
        literature = openpyxl.load_workbook(job / "export" / "literature.xlsx")
        self.assertEqual(literature.sheetnames, ["文献表"])
        sheet = literature["文献表"]
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(header, export.LITERATURE_COLUMNS)
        self.assertEqual(sheet.max_row, 4)  # 表头 + 3 行
        self.assertEqual(sheet.cell(1, 1).value, "paper_id")
        parameters = openpyxl.load_workbook(job / "export" / "parameters.xlsx")
        self.assertEqual(parameters.sheetnames, ["参数明细"])
        sheet = parameters["参数明细"]
        self.assertEqual([cell.value for cell in sheet[1]], export.PARAMETER_COLUMNS)
        self.assertEqual(sheet.max_row, 3)  # 表头 + 2 行（阻断行被剔除）
        values = {sheet.cell(index, 1).value: sheet.cell(index, 13).value for index in range(2, sheet.max_row + 1)}
        self.assertEqual(values.get("GSC-002-R001"), "pending_review")
        self.assertEqual(sheet.cell(2, 6).value, "maize")   # 文本字段原样
        self.assertEqual(sheet.cell(3, 6).value, 6.5)       # 数值字段保持数字类型

    def test_xlsx_is_a_valid_zip_package(self):
        job = self.default_job(name="zip")
        self.export(job)
        for name in ("literature.xlsx", "parameters.xlsx"):
            with zipfile.ZipFile(job / "export" / name) as archive:
                self.assertIn("xl/workbook.xml", archive.namelist())
                workbook = archive.read("xl/workbook.xml").decode("utf-8")
                self.assertIn("文献表" if name.startswith("literature") else "参数明细", workbook)

    # ------------------------------------------------------------------ versioning

    def test_reexport_archives_previous_version(self):
        job = self.default_job(name="archive")
        self.export(job)
        first = (job / "export" / "literature.csv").read_bytes()
        code, events = self.export(job)
        self.assertEqual(code, 0)
        archived = events[-1]["counts"]  # counts 不含归档信息，改从返回值/报告读取
        report = (job / "export" / "export_report.md").read_text(encoding="utf-8")
        self.assertIn("旧版本已归档：export/history/", report)
        history_root = job / "export" / "history"
        stamps = sorted(path for path in history_root.iterdir() if path.is_dir())
        self.assertEqual(len(stamps), 1)
        archived_csv = stamps[0] / "literature.csv"
        self.assertTrue(archived_csv.is_file())
        self.assertEqual(archived_csv.read_bytes(), first)
        self.assertEqual(sorted(path.name for path in stamps[0].iterdir()),
                         sorted(Path(name).name for name in export.EXPORT_FILES))
        current = (job / "export" / "literature.csv").read_bytes()
        self.assertEqual(current, first)
        _ = archived

    def test_archive_previous_returns_none_when_empty(self):
        empty = self.root / "empty-export"
        empty.mkdir()
        self.assertIsNone(export.archive_previous(empty))
        self.assertIsNone(export.archive_previous(self.root / "missing-export"))

    # ------------------------------------------------------------------ templates

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl 未安装")
    def test_user_template_single_header_mapping(self):
        import openpyxl  # noqa: PLC0415
        template = self.root / "template-single.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "文献表"
        sheet.append(["paper_id", "title", "题名备注"])
        sheet.append(["示例", "示例题名", "示例备注"])
        workbook.save(template)

        job = self.default_job(name="template-single")
        config = json.loads((job / "run_config.json").read_text(encoding="utf-8"))
        config["templates"] = {"literature": str(template), "parameters": None}
        (job / "run_config.json").write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        code, events = self.export(job)
        self.assertEqual(code, 0)
        report = (job / "export" / "export_report.md").read_text(encoding="utf-8")
        self.assertIn("## 用户模板映射（供研究者确认）", report)
        self.assertIn("未识别的模板表头：题名备注", report)
        self.assertIn("模板未覆盖的表列", report)
        sheet = openpyxl.load_workbook(job / "export" / "literature.xlsx")["文献表"]
        self.assertEqual(sheet.cell(2, 1).value, "GSC-001")
        self.assertEqual(sheet.cell(1, 1).value, "paper_id")
        self.assertEqual(sheet.max_row, 4)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl 未安装")
    def test_user_template_two_row_header_mapping(self):
        import openpyxl  # noqa: PLC0415
        template = self.root / "template-double.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "参数明细"
        sheet["A1"] = "参数"
        sheet.merge_cells("A1:B1")
        sheet["C1"] = "出处"
        sheet["A2"] = "row_id"
        sheet["B2"] = "field"
        sheet["C2"] = "evidence_quote"
        workbook.save(template)

        job = self.default_job(name="template-double")
        config = json.loads((job / "run_config.json").read_text(encoding="utf-8"))
        config["templates"] = {"literature": None, "parameters": str(template)}
        (job / "run_config.json").write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        code, _ = self.export(job)
        self.assertEqual(code, 0)
        sheet = openpyxl.load_workbook(job / "export" / "parameters.xlsx")["参数明细"]
        self.assertEqual(sheet.cell(1, 1).value, "参数")
        self.assertEqual(sheet.cell(2, 1).value, "row_id")
        self.assertEqual(sheet.cell(3, 1).value, "GSC-001-R002")
        self.assertEqual(sheet.cell(3, 2).value, "crop")
        self.assertEqual(sheet.cell(3, 3).value, "continuous maize")
        report = (job / "export" / "export_report.md").read_text(encoding="utf-8")
        self.assertIn("表头行数：2", report)

    def test_missing_template_errors(self):
        job = self.default_job(name="template-missing")
        config = json.loads((job / "run_config.json").read_text(encoding="utf-8"))
        config["templates"] = {"literature": "templates/does-not-exist.xlsx", "parameters": None}
        (job / "run_config.json").write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        code, events = self.export(job)
        self.assertEqual(code, 1)
        self.assertIn("用户模板不可用", events[-1]["message"])

    # ------------------------------------------------------------------ errors

    def test_literature_rows_fall_back_to_candidates_without_dedupe(self):
        job = self.build_job("nodedupe", groups=None)
        code, events = self.export(job)
        self.assertEqual(code, 0)
        with (job / "export" / "literature.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([entry["paper_id"] for entry in rows], ["GSC-001", "GSC-B2-003", "GSC-002"])
        first = rows[0]
        self.assertEqual(first["screening_status"], "pending_review")
        self.assertEqual(first["fulltext_status"], "fetched_html")
        duplicate = rows[1]
        self.assertEqual(duplicate["screening_status"], "auto_excluded_duplicate")
        self.assertEqual(duplicate["fulltext_status"], "")

    def test_literature_rows_drop_blocked_ids(self):
        records = base_records() + [{"candidate_id": "", "title": "无 ID 记录", "screening_status": "pending_review"}]
        groups = [base_groups()[0],
                  {**base_groups()[1], "records": [base_records()[2], {"candidate_id": "", "title": "无 ID 记录"}]}]
        job = self.build_job("blockedlit", records=records, groups=groups)
        validation = json.loads((job / "validation" / "validation.json").read_text(encoding="utf-8"))
        self.assertIn("invalid_paper_id", [rule for issue in validation["blocking"] for rule in issue["rules"]])
        code, _ = self.export(job)
        self.assertEqual(code, 0)
        with (job / "export" / "literature.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertNotIn("", [entry["paper_id"] for entry in rows])
        self.assertEqual(len(rows), 3)

    def test_export_requires_validation(self):
        job = self.build_job("novalid")
        shutil.rmtree(job / "validation")
        code, events = self.export(job)
        self.assertEqual(code, 1)
        self.assertEqual(events[-1]["message"], "尚未完成校验（缺少 validation/validation.json）")

    def test_export_without_confirmations_still_exports(self):
        job = self.build_job("noschema")
        (job / "confirmations.json").unlink()
        code, events = self.export(job)
        self.assertEqual(code, 0)
        report = (job / "export" / "export_report.md").read_text(encoding="utf-8")
        self.assertIn("（未确认字段定义）", report)

    def test_safe_csv_cell_blocks_formula_injection(self):
        docs = base_docs()
        docs["GSC-001"] = copy.deepcopy(docs["GSC-001"])
        docs["GSC-001"]["rows"][0]["value_raw"] = "=cmd|'/c calc'!A1"
        job = self.build_job("formula", docs=docs)
        self.export(job)
        text = (job / "export" / "parameters.csv").read_text(encoding="utf-8-sig")
        self.assertIn("'=cmd", text)

    def test_map_sheet_helpers(self):
        row_data = row()
        self.assertEqual(export._cell(None), "")
        self.assertEqual(export._cell({"b": 1, "a": 2}), '{"a":2,"b":1}')
        self.assertEqual(export._cell(0), 0)
        self.assertEqual(export.normalize_header("Paper ID"), "paperid")
        self.assertEqual(export.normalize_header("论文_ID"), "论文id")
        _ = row_data


if __name__ == "__main__":
    unittest.main()
