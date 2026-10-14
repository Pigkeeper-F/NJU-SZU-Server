from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


PROVINCES = [
    "北京", "天津", "河北", "山西", "内蒙古", "辽宁", "吉林", "黑龙江", "上海", "江苏",
    "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南", "广东", "广西",
    "海南", "重庆", "四川", "贵州", "云南", "西藏", "陕西", "甘肃", "青海", "宁夏",
    "新疆", "台湾", "香港", "澳门",
]

PROVINCE_ALIASES = {
    "Beijing": "北京",
    "Tianjin": "天津",
    "Hebei": "河北",
    "Shanxi": "山西",
    "Inner Mongolia": "内蒙古",
    "Liaoning": "辽宁",
    "Jilin": "吉林",
    "Heilongjiang": "黑龙江",
    "Shanghai": "上海",
    "Jiangsu": "江苏",
    "Zhejiang": "浙江",
    "Anhui": "安徽",
    "Fujian": "福建",
    "Jiangxi": "江西",
    "Shandong": "山东",
    "Henan": "河南",
    "Hubei": "湖北",
    "Hunan": "湖南",
    "Guangdong": "广东",
    "Guangxi": "广西",
    "Hainan": "海南",
    "Chongqing": "重庆",
    "Sichuan": "四川",
    "Guizhou": "贵州",
    "Yunnan": "云南",
    "Tibet": "西藏",
    "Shaanxi": "陕西",
    "Gansu": "甘肃",
    "Qinghai": "青海",
    "Ningxia": "宁夏",
    "Xinjiang": "新疆",
}


def norm_id(value) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return str(int(float(text)))
    except Exception:
        return text


def header_map(ws) -> dict[str, int]:
    mapping = {}
    for col in range(1, ws.max_column + 1):
        group = ws.cell(1, col).value
        field = ws.cell(2, col).value
        if field in (None, ""):
            if group in (None, ""):
                continue
            field = str(group).strip()
            group = ""
        else:
            field = str(field).strip()
        group = str(group).strip() if group not in (None, "") else ""
        mapping[field] = col
        if group:
            mapping[f"{group}.{field}"] = col
    return mapping


def extract_year(value) -> int | None:
    if value in (None, ""):
        return None
    text = str(value)
    years = [int(y) for y in re.findall(r"(?:19|20)\d{2}", text)]
    years = [y for y in years if 1900 <= y <= 2026]
    if not years:
        return None
    return years[0]


def decade(year: int | None) -> str:
    if year is None:
        return "未注明"
    return f"{year // 10 * 10}s"


def normalize_depth(value, row, cols) -> str:
    text = "" if value in (None, "") else str(value).strip()
    if text:
        text = text.replace("—", "-").replace("–", "-").replace("－", "-")
        text = re.sub(r"\s+", "", text)
        return text
    # Fallback from layer-specific SOC columns when depth text is missing.
    for key, label in [
        ("SOC.SOC     0-30cm", "0-30"),
        ("SOC.SOC     30-60cm", "30-60"),
        ("SOC.SOC     60-100cm", "60-100"),
    ]:
        col = cols.get(key)
        if col and row[col - 1] not in (None, ""):
            return label
    return "未注明"


def depth_band(depth: str) -> str:
    nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", depth)]
    if len(nums) >= 2:
        top, bottom = nums[0], nums[1]
        if top == 0 and bottom <= 30:
            return "表层 0-30 cm"
        if bottom <= 30:
            return "表层 0-30 cm"
        if top < 30 < bottom:
            return "跨 30 cm"
        if 30 <= top and bottom <= 60:
            return "30-60 cm"
        if 60 <= top and bottom <= 100:
            return "60-100 cm"
        if bottom > 100:
            return ">100 cm"
        return "其他土层"
    if depth == "未注明":
        return "未注明"
    return "其他/无法解析"


def get_province(region) -> str:
    text = "" if region in (None, "") else str(region)
    for province in PROVINCES:
        if province in text:
            return province
    lower = text.lower()
    for alias, province in PROVINCE_ALIASES.items():
        if alias.lower() in lower:
            return province
    return "未识别"


def counter_to_list(counter: Counter, limit: int | None = None):
    items = counter.most_common(limit)
    return [{"name": k, "count": v} for k, v in items]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(Path(args.workbook), read_only=True, data_only=True)
    ws = wb["点位试验"]
    cols = header_map(ws)

    no_col = cols.get("No.", 1)
    year_col = cols.get("采样年份", 4)
    region_col = cols.get("点位信息.Region", cols.get("Region", 5))
    lon_col = cols.get("点位信息.经度", 6)
    lat_col = cols.get("点位信息.纬度", 7)
    crop_col = cols.get("作物信息.Crop type", cols.get("Crop type", 17))
    soc_col = cols.get("土壤信息.SOC (g/kg)", cols.get("SOC (g/kg)", 16))
    depth_col = cols.get("SOC.土层（cm）", 55)

    rows = []
    unique_ids = set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        no = norm_id(row[no_col - 1])
        if not no:
            continue
        unique_ids.add(no)
        year = extract_year(row[year_col - 1])
        depth = normalize_depth(row[depth_col - 1] if depth_col else None, row, cols)
        lon = row[lon_col - 1] if lon_col else None
        lat = row[lat_col - 1] if lat_col else None
        region = row[region_col - 1] if region_col else None
        crop = row[crop_col - 1] if crop_col else None
        rows.append(
            {
                "no": no,
                "year": year,
                "depth": depth,
                "depth_band": depth_band(depth),
                "province": get_province(region),
                "has_coord": lon not in (None, "") and lat not in (None, ""),
                "crop": "未注明" if crop in (None, "") else str(crop).strip(),
                "has_soc": any(
                    row[col - 1] not in (None, "")
                    for col in (soc_col, 45, 48, 51)
                    if col and col <= len(row)
                ),
            }
        )

    by_depth = Counter(r["depth"] for r in rows)
    by_depth_band = Counter(r["depth_band"] for r in rows)
    by_year = Counter(r["year"] if r["year"] is not None else "未注明" for r in rows)
    by_decade = Counter(decade(r["year"]) for r in rows)
    by_province = Counter(r["province"] for r in rows)

    ids_by_year = defaultdict(set)
    ids_by_decade = defaultdict(set)
    ids_by_depth_band = defaultdict(set)
    for r in rows:
        ids_by_year[r["year"] if r["year"] is not None else "未注明"].add(r["no"])
        ids_by_decade[decade(r["year"])].add(r["no"])
        ids_by_depth_band[r["depth_band"]].add(r["no"])

    year_rows = [
        {"year": k, "rows": by_year[k], "articles": len(ids_by_year[k])}
        for k in sorted(by_year, key=lambda x: 9999 if x == "未注明" else int(x))
    ]
    decade_rows = [
        {"period": k, "rows": by_decade[k], "articles": len(ids_by_decade[k])}
        for k in sorted(by_decade, key=lambda x: 9999 if x == "未注明" else int(str(x)[:4]))
    ]
    depth_band_rows = [
        {"depth_band": k, "rows": by_depth_band[k], "articles": len(ids_by_depth_band[k])}
        for k in by_depth_band
    ]

    summary = {
        "workbook": str(Path(args.workbook)),
        "sheet": "点位试验",
        "total_rows": len(rows),
        "unique_articles": len(unique_ids),
        "rows_with_soc": sum(1 for r in rows if r["has_soc"]),
        "rows_with_coordinates": sum(1 for r in rows if r["has_coord"]),
        "coordinate_coverage_percent": round(sum(1 for r in rows if r["has_coord"]) / len(rows) * 100, 1) if rows else 0,
        "year_min": min((r["year"] for r in rows if r["year"] is not None), default=None),
        "year_max": max((r["year"] for r in rows if r["year"] is not None), default=None),
        "missing_year_rows": sum(1 for r in rows if r["year"] is None),
        "depth_band_distribution": depth_band_rows,
        "top_exact_depths": counter_to_list(by_depth, 20),
        "decade_distribution": decade_rows,
        "year_distribution": year_rows,
        "top_provinces": counter_to_list(by_province, 15),
        "top_crops": counter_to_list(Counter(r["crop"] for r in rows), 15),
    }

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
