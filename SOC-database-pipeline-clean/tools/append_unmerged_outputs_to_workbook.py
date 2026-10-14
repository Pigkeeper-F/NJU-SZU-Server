from __future__ import annotations

import argparse
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl


SHEETS = ("点位试验", "长期定位试验")


def composite_key(field_entry: dict) -> str:
    group = field_entry.get("group")
    field = field_entry["field"]
    return f"{group}.{field}" if group else field


def norm_id(value) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return str(int(float(text)))
    except Exception:
        return text


def existing_ids(wb) -> set[str]:
    ids: set[str] = set()
    for sheet_name in SHEETS:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=3, values_only=True):
            no = norm_id(row[0] if row else None)
            if no:
                ids.add(no)
    return ids


def last_data_row(ws) -> int:
    for row in range(ws.max_row, 2, -1):
        if ws.cell(row, 1).value not in (None, ""):
            return row
    return 2


def copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    if source_row < 1:
        return
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def write_json_row(ws, schema: list[dict], target_row: int, data: dict, source_id: str) -> None:
    for field in schema:
        key = composite_key(field)
        value = source_id if key == "No." else data.get(key)
        if value not in (None, ""):
            ws.cell(target_row, field["col"], value)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--repo", required=True)
    parser.add_argument("--report", required=True)
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    repo = Path(args.repo)
    report_path = Path(args.report)
    schema_dir = repo / "schemas"
    output_dir = repo / "outputs"

    schemas = {
        sheet: json.loads((schema_dir / f"{sheet}.json").read_text(encoding="utf-8"))
        for sheet in SHEETS
    }

    wb = openpyxl.load_workbook(workbook_path)
    ids_before = existing_ids(wb)

    candidates = []
    skipped = []
    for path in sorted(output_dir.glob("*.json"), key=lambda p: int(p.stem) if p.stem.isdigit() else 10**9):
        data = json.loads(path.read_text(encoding="utf-8"))
        if data.get("status") == "skipped":
            continue
        sid = path.stem
        if sid in ids_before:
            continue
        sheet = data.get("sheet")
        if sheet not in SHEETS:
            skipped.append((sid, f"invalid sheet: {sheet!r}"))
            continue
        candidates.append((sid, path, data))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = workbook_path.with_name(f"{workbook_path.stem}.bak_before_append_outputs_{timestamp}{workbook_path.suffix}")
    shutil.copy2(workbook_path, backup_path)

    appended_files = []
    appended_rows_by_sheet = {sheet: 0 for sheet in SHEETS}
    append_start_rows = {}

    for sid, path, data in candidates:
        sheet = data["sheet"]
        ws = wb[sheet]
        schema = schemas[sheet]
        source_style_row = last_data_row(ws)
        next_row = source_style_row + 1
        append_start_rows.setdefault(sheet, next_row)
        for row_data in data.get("rows") or []:
            copy_row_style(ws, source_style_row, next_row, max(f["col"] for f in schema))
            write_json_row(ws, schema, next_row, row_data, sid)
            next_row += 1
            appended_rows_by_sheet[sheet] += 1
        appended_files.append((sid, sheet, len(data.get("rows") or []), path.name))

    wb.save(workbook_path)

    wb_check = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ids_after = existing_ids(wb_check)
    missing_after = [
        path.stem
        for path in output_dir.glob("*.json")
        if path.stem not in ids_after
        and json.loads(path.read_text(encoding="utf-8")).get("status") != "skipped"
    ]

    lines = [
        "# Append unmerged outputs report",
        "",
        f"Workbook: {workbook_path}",
        f"Backup: {backup_path}",
        f"Outputs considered missing before append: {len(candidates)}",
        f"Appended rows by sheet: {appended_rows_by_sheet}",
        f"Missing non-skipped outputs after append: {len(missing_after)}",
        "",
        "## Appended files",
    ]
    for sid, sheet, rows, name in appended_files:
        lines.append(f"- {sid}: {sheet}, {rows} rows, {name}")
    if skipped:
        lines.extend(["", "## Skipped"])
        for sid, reason in skipped:
            lines.append(f"- {sid}: {reason}")
    if missing_after:
        lines.extend(["", "## Missing after append"])
        lines.append(", ".join(sorted(missing_after, key=lambda x: int(x) if x.isdigit() else 10**9)))
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"backup={backup_path}")
    print(f"appended_files={len(appended_files)}")
    print(f"appended_rows_by_sheet={appended_rows_by_sheet}")
    print(f"missing_after={len(missing_after)}")
    print(f"report={report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
