from __future__ import annotations

import sys
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import merge


class MergeTemplateCleanupTest(unittest.TestCase):
    def test_template_data_rows_are_all_cleared_before_merge(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["header"])
        ws.append(["field"])
        ws.append(["old row 1"])
        ws.append(["old row 2"])
        ws.append(["old row 3"])

        merge.clear_sample_row(ws)

        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(1, 1).value, "header")
        self.assertEqual(ws.cell(2, 1).value, "field")


if __name__ == "__main__":
    unittest.main()
